"""
Integration between existing scrapers and Clay-like enrichment features
"""

from typing import Dict, List, Optional, Any
from clay_like_enrichment import ClayLikePlatform, DataEnrichmentEngine
from google_maps_scraper import GoogleMapsScraper, scrape_business_info
from google_web_scraper import GoogleWebScraper, search_businesses_web
from social_media_search_scraper import SocialMediaSearchScraper


class IntegratedScraper:
    """
    Integrated scraper that combines Google Maps/Web scraping with Clay-like enrichment
    """
    
    def __init__(self):
        self.platform = ClayLikePlatform()
        self._setup_enrichment_sources()
    
    def _setup_enrichment_sources(self):
        """Setup enrichment sources from existing scrapers."""
        
        # Google Maps as enrichment source
        def enrich_from_maps(record: Dict) -> Dict:
            """Enrich record using Google Maps."""
            if not record.get('name'):
                return {}
            
            try:
                location = record.get('location') or record.get('address', '')
                results = scrape_business_info(
                    business_name=record['name'],
                    location=location,
                    headless=True,
                    export_excel=False,
                    export_leads=False
                )
                
                if results:
                    return results[0]  # Return first result
            except:
                pass
            return {}
        
        # Google Web as enrichment source
        def enrich_from_web(record: Dict) -> Dict:
            """Enrich record using Google Web search."""
            if not record.get('name'):
                return {}
            
            try:
                query = f"{record['name']} {record.get('location', '')}"
                results = search_businesses_web(
                    query=query,
                    max_results=1,
                    max_pages=1,
                    headless=True,
                    export_excel=False,
                    export_csv=False
                )
                
                if results:
                    return results[0]
            except:
                pass
            return {}
        
        # Social Media as enrichment source
        def enrich_from_social(record: Dict) -> Dict:
            """Enrich record using Social Media search."""
            if not record.get('name'):
                return {}
            
            try:
                scraper = SocialMediaSearchScraper(headless=True)
                try:
                    # Try Facebook first
                    fb_results = scraper.search_facebook(record['name'], max_results=1)
                    if fb_results:
                        social_data = {
                            'facebook_url': fb_results[0].get('url'),
                            'facebook_followers': fb_results[0].get('followers'),
                            'facebook_likes': fb_results[0].get('likes'),
                            'facebook_category': fb_results[0].get('category')
                        }
                        return social_data
                finally:
                    scraper.close()
            except:
                pass
            return {}
        
        # Add enrichment sources
        self.platform.enrichment_engine.add_enrichment_source("Google Maps", enrich_from_maps)
        self.platform.enrichment_engine.add_enrichment_source("Google Web", enrich_from_web)
        self.platform.enrichment_engine.add_enrichment_source("Social Media", enrich_from_social)
    
    def enrich_lead(self, lead: Dict, fields: List[str] = None) -> Dict:
        """
        Enrich a lead using waterfall enrichment across multiple sources.
        
        Args:
            lead: Lead dictionary with at least 'name' field
            fields: List of fields to enrich (default: common fields)
        
        Returns:
            Enriched lead dictionary
        """
        if fields is None:
            fields = ['email', 'phone', 'website', 'address', 'rating', 'category', 
                     'facebook_url', 'instagram_url', 'linkedin_url']
        
        print(f"Enriching lead: {lead.get('name', 'Unknown')}")
        enriched = self.platform.enrichment_engine.waterfall_enrich(lead, fields)
        
        # Calculate lead score
        enriched['lead_score'] = self._calculate_lead_score(enriched)
        enriched['lead_status'] = self._determine_lead_status(enriched['lead_score'])
        
        return enriched
    
    def enrich_bulk(self, leads: List[Dict], fields: List[str] = None) -> List[Dict]:
        """Enrich multiple leads."""
        enriched_leads = []
        for i, lead in enumerate(leads, 1):
            print(f"\n[{i}/{len(leads)}] Processing lead...")
            enriched = self.enrich_lead(lead, fields)
            enriched_leads.append(enriched)
        return enriched_leads
    
    def _calculate_lead_score(self, lead: Dict) -> int:
        """Calculate lead score based on data completeness."""
        score = 0
        if lead.get('name'):
            score += 10
        if lead.get('email'):
            score += 25
        if lead.get('phone'):
            score += 20
        if lead.get('website'):
            score += 15
        if lead.get('address'):
            score += 15
        if lead.get('rating'):
            score += 10
        if lead.get('facebook_url') or lead.get('instagram_url') or lead.get('linkedin_url'):
            score += 5
        
        return min(score, 100)
    
    def _determine_lead_status(self, score: int) -> str:
        """Determine lead status based on score."""
        if score >= 70:
            return "Hot"
        elif score >= 50:
            return "Warm"
        elif score >= 30:
            return "Cold"
        else:
            return "New"
    
    def create_enrichment_workflow(self, name: str = "Standard Lead Enrichment") -> Dict:
        """Create a standard enrichment workflow."""
        workflow = self.platform.create_workflow(name, [
            self.platform.workflow_automation.add_conditional_step(
                condition="email == None",
                action_if_true=self.platform.workflow_automation.add_enrichment_step(
                    ['email', 'phone', 'website']
                ),
                action_if_false=self.platform.workflow_automation.add_enrichment_step(
                    ['phone', 'website']
                )
            ),
            self.platform.workflow_automation.add_conditional_step(
                condition="lead_score > 50",
                action_if_true=self.platform.workflow_automation.add_action_step(
                    action_type='sync_crm',
                    destination='salesforce'
                )
            )
        ])
        return workflow
    
    def execute_enrichment_workflow(self, workflow_name: str, lead: Dict) -> Dict:
        """Execute enrichment workflow on a lead."""
        # First enrich the lead
        enriched = self.enrich_lead(lead)
        
        # Then execute workflow
        result = self.platform.execute_workflow(workflow_name, enriched)
        return result
    
    def build_high_intent_audience(self, leads: List[Dict]) -> List[Dict]:
        """Build audience of high-intent leads."""
        audience = self.platform.build_audience("High Intent Leads", {
            'lead_score': {'operator': 'greater_than', 'value': 60},
            'email': {'operator': 'is_not_empty', 'value': None},
            'phone': {'operator': 'is_not_empty', 'value': None}
        })
        
        matched = self.platform.audience_builder.match_records("High Intent Leads", leads)
        return matched
    
    def sync_to_crm(self, crm_type: str, leads: List[Dict]):
        """Sync enriched leads to CRM."""
        return self.platform.sync_to_crm(crm_type, leads)
    
    def create_outbound_sequence(self, name: str = "Cold Outreach") -> Dict:
        """Create an outbound email sequence."""
        sequence = self.platform.create_email_sequence(name, [
            self.platform.email_sequencer.add_email_step(
                subject="Quick question about {{company}}",
                body_template="""Hi {{name}},

I noticed {{company}} and thought you might be interested in...

Best regards""",
                delay_days=0
            ),
            self.platform.email_sequencer.add_wait_step(3),
            self.platform.email_sequencer.add_email_step(
                subject="Following up on {{company}}",
                body_template="""Hi {{name}},

Just following up on my previous email about {{company}}...

Best regards""",
                delay_days=0
            ),
            self.platform.email_sequencer.add_wait_step(4),
            self.platform.email_sequencer.add_email_step(
                subject="Last attempt - {{company}}",
                body_template="""Hi {{name}},

This will be my last email, but I wanted to make sure you saw this opportunity for {{company}}...

Best regards""",
                delay_days=0
            )
        ])
        return sequence


def export_enriched_leads_to_excel(leads: List[Dict], filename: str = "enriched_leads.xlsx", 
                                   output_folder: str = "excel_results"):
    """Export enriched leads to Excel with all Clay-like fields."""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required. Install it with: pip install openpyxl")
    
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    
    filepath = os.path.join(output_folder, filename)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enriched Leads"
    
    # Headers with all enrichment fields
    headers = [
        "Name", "Email", "Phone", "Website", "Address", "Company",
        "Rating", "Reviews", "Category", "Business Hours",
        "Facebook URL", "Instagram URL", "LinkedIn URL", "Twitter URL",
        "Facebook Followers", "Instagram Followers",
        "Lead Score", "Lead Status", "Data Completeness",
        "Enrichment Sources", "Last Enriched"
    ]
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # Write data
    for row_num, lead in enumerate(leads, 2):
        ws.cell(row=row_num, column=1, value=lead.get('name', ''))
        ws.cell(row=row_num, column=2, value=lead.get('email', ''))
        ws.cell(row=row_num, column=3, value=lead.get('phone', ''))
        ws.cell(row=row_num, column=4, value=lead.get('website', ''))
        ws.cell(row=row_num, column=5, value=lead.get('address', ''))
        ws.cell(row=row_num, column=6, value=lead.get('company', ''))
        ws.cell(row=row_num, column=7, value=lead.get('rating', ''))
        ws.cell(row=row_num, column=8, value=lead.get('reviews_count', ''))
        ws.cell(row=row_num, column=9, value=lead.get('category', ''))
        ws.cell(row=row_num, column=10, value=lead.get('business_hours', ''))
        ws.cell(row=row_num, column=11, value=lead.get('facebook', lead.get('facebook_url', '')))
        ws.cell(row=row_num, column=12, value=lead.get('instagram', lead.get('instagram_url', '')))
        ws.cell(row=row_num, column=13, value=lead.get('linkedin', lead.get('linkedin_url', '')))
        ws.cell(row=row_num, column=14, value=lead.get('twitter', ''))
        ws.cell(row=row_num, column=15, value=lead.get('facebook_followers', ''))
        ws.cell(row=row_num, column=16, value=lead.get('instagram_followers', ''))
        ws.cell(row=row_num, column=17, value=lead.get('lead_score', 0))
        ws.cell(row=row_num, column=18, value=lead.get('lead_status', 'New'))
        
        # Calculate data completeness
        filled_fields = sum(1 for key in ['name', 'email', 'phone', 'website', 'address'] if lead.get(key))
        completeness = f"{(filled_fields/5)*100:.1f}%"
        ws.cell(row=row_num, column=19, value=completeness)
        
        ws.cell(row=row_num, column=20, value=lead.get('enrichment_sources', ''))
        ws.cell(row=row_num, column=21, value=lead.get('last_enriched', ''))
    
    # Auto-adjust column widths
    for col_num, header in enumerate(headers, 1):
        column_letter = get_column_letter(col_num)
        max_length = len(header)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_num, max_col=col_num):
            if row[0].value:
                max_length = max(max_length, len(str(row[0].value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    wb.save(filepath)
    print(f"\nEnriched leads exported to {filepath}")
    return filepath

